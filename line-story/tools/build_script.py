#!/usr/bin/env python3
"""Build the printable teacher script from the exact classroom cue data.
Usage: python tools/build_script.py
Requires: python-docx (document build only; classroom has no dependencies).
"""
from __future__ import annotations
import hashlib, json, math, re, zipfile
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'scripts/lesson-data.js'
def load_lesson():
    text=DATA.read_text(encoding='utf-8')
    return json.loads(text.split('window.LINE_LESSON = ',1)[1].rstrip().removesuffix(';'))
L=load_lesson(); C=L['cues']; SHA=hashlib.sha256(DATA.read_bytes()).hexdigest()
DOC=ROOT/'teacher_script.docx'; MD=ROOT/'teacher_script.md'
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.7)
INK='000000'; GOLD='000000'; MUTED='4B5553'
def style_font(st,size,name='Noto Serif CJK SC',bold=False,color=INK):
    st.font.name=name;st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor.from_string(color)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
style_font(doc.styles['Normal'],11)
normal=doc.styles['Normal'].paragraph_format
normal.line_spacing=1.28;normal.space_after=Pt(5)
for name,size in [('Title',29),('Heading 1',19),('Heading 2',14),('Heading 3',11)]:
    style_font(doc.styles[name],size,'Noto Sans CJK SC',True)
    pf=doc.styles[name].paragraph_format;pf.space_before=Pt(10);pf.space_after=Pt(8);pf.keep_with_next=True
style_font(doc.styles['Caption'],9,'Noto Sans CJK SC',False,MUTED)
for name in ['Header','Footer']:
    style_font(doc.styles[name],8.5,'Noto Sans CJK SC',False,MUTED)
# Remove template paragraph borders, including Word's blue Title underline.
for st in doc.styles:
    for border in st.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
header=sec.header.paragraphs[0];header.text='《童年》整本书阅读  /  线条叙事版教师用稿'
p=sec.footer.paragraphs[0];p.alignment=2
p.add_run('线条叙事 · ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
doc.core_properties.title='童年公开课教师逐字稿'
doc.core_properties.subject=f"40分钟公开课；与{len(C)}个网页节拍和{len(L['nodes'])}个授课节点同步"
doc.core_properties.author='';doc.core_properties.last_modified_by=''
md=['# 童年公开课教师逐字稿','',L['title'],'',f"{L['version']} · 40分钟公开课 · {len(C)}个精细节拍 · {len(L['nodes'])}个授课节点",'',f'课堂数据 SHA-256：`{SHA}`','']
def text(s,style=None):
    p=doc.add_paragraph(s,style);p.paragraph_format.widow_control=True;return p
def label(name,s):
    if not s:return
    p=doc.add_paragraph();r=p.add_run(name+'  ');r.bold=True;r.font.name='Noto Sans CJK SC';r.font.color.rgb=RGBColor.from_string(GOLD)
    p.add_run(s);p.paragraph_format.space_after=Pt(5);return p
def page():doc.add_page_break()
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths:
        for i,width in enumerate(widths):t.columns[i].width=Cm(width)
    tp=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    tp.append(borders)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();nc=OxmlElement('w:cantSplit');pr.append(nc)
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            if widths:cell.width=Cm(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8ECEE' if ri==0 else ('F5F7F8' if ri%2 else 'FFFFFF'));cp.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.2
                for r in p.runs:r.font.size=Pt(9);r.bold=ri==0
    return t
# Opening pages, intentionally separate from live spoken script.
text('六年级 · 整本书阅读', 'Caption')
text('童年公开课教师逐字稿','Title')
text('童年里什么值得被留下','Subtitle')
style_font(doc.styles['Subtitle'],16,'Noto Sans CJK SC',False,INK)
text('从“哪一种影响更大”，走向“什么样的生活值得被守住”。')
text(f"本稿与网页{len(C)}个精细节拍逐一对应，并归入{len(L['nodes'])}个教师可辨认的授课节点。教师话语、画面动作、停留和节点时间均从同一份课堂数据生成。主线依次完成共同阅读、三轮个人阅读游戏、一次小组讨论与个人书写；教师以课堂反应掌握推进节奏。")
text('长段原文保持阅读主体。先让学生说出证据，再由教师推进关系变化；象征笔迹只表示表达与修订发生，不代写答案。窗边长段与安妮的三处独立短引均进入主线。')
doc.add_heading('这节课要保留的三层认识',2)
label('看见影响','伤害的深度与支持的力量并非同一种作用。允许学生保留不同判断，但要求说明尺度、给出证据。')
label('检验理由','理解一个人的经历，不等于认可他伤害别人。概括要补回代价，推论要经原文检验，赞美要落到人的选择。')
label('守住日常','爱与帮助真实可贵，但不自动消除伤害和限制。幸福还需要安全、学习、稳定生活与被尊重等条件得到持续保障。')
doc.add_heading('课堂准备与边界',2)
text('学生准备同版《童年》、B5阅读本和笔；不需打印任务纸。活动依次为全班证据交流、三轮个人阅读游戏、一次小组讨论和退出卡。未通读的学生可从共同文本进入。')
text('不组织殴打表演，不用流血或灾难画面刺激学生；不要求公开家庭创伤。不以“谁更苦”比较孩子，也不把家庭暴力、纳粹迫害和战争造成的困境解释成同一原因。')
text('建议时间不是强制倒计时。教师依据课堂掌握推进节奏；超时先减少复述和展示人数，不削去两分钟并读、五分钟讨论与九十秒独立书写。')
text(L['version']+' · 2026年10月24日','Caption')
page();doc.add_heading('授课路线与操作',1)
table(['节点','建议时间','线的变化'],[(f"{n['number']:02d} {n['title']}",n['timing'],n['line']) for n in L['nodes']],[6.0,3.5,7.5])
md.extend(['## 十五个授课节点','','| 节点 | 建议时间 | 线的变化 |','| --- | --- | --- |'])
md.extend(f"| {n['number']:02d} {n['title']} | {n['timing']} | {n['line']} |" for n in L['nodes'])
md.append('')
page();doc.add_heading('操作与课前演练',1)
operations=[('下一步／上一步','→、空格、PageDown／←、PageUp'),('全屏／目录','F／M'),('打开独立教师台','P 或 N'),('暂停／继续活动计时','T'),('只重播当前动作','R，不重置活动计时'),('重置当前活动计时','Shift + R'),('暂停／继续画面动作','A'),('暂隐学生画面','B；再按 B 或 Esc 恢复'),('前一／后一授课节点','[ ／ ]'),('关闭回读或目录','Esc')]
table(['操作','按键'],operations,[7.2,9.8])
operation_notes=[
 '每次下一步进入一个有限动作，结束后画面停住。动作未完时，第一次推进只完成当前动作，再按一次才进入下一节拍。返回和目录跳转直接恢复目标状态；主动重播只重演当前动作。',
 '默认空白点击和右键都不推进，使用屏幕按钮或翻页笔即可。需要右键推进时，在教师台明确开启。计时结束只提醒，不自动翻页；阅读与学生回答可以一直停留。',
 '教师台是独立窗口，显示当前旁白、下一节拍、活动计时与节点导航。使用系统“扩展显示”，把学生窗口放到投影屏，教师台留在电脑屏；镜像显示不能隐藏教师台。若弹窗被拦截，允许本页弹出窗口。',
 '三张真实图片在指定节拍展开、后续节拍收回；由教师按键配合旁白，不使用麦克风或语音识别。列宾原画完整铺展，不裁切人物、不叠标题；两份纪实照片各留38秒，含讲述后的安静观看。',
 '教师不在网页输入答案。学生纸笔书写；原文词语可点亮、箭头可检验，讨论中可回看已读材料。小组笔迹只表示已有表达与修订，不冒充实际发言。',
 '课前用实际投影和翻页笔检查长段原文、全屏画作、保护问答、教师台与最后45秒。刷新后恢复当前会话进度；“重新开始”会清除本课记录。默认无音乐、外部脚本和网络字体。',
]
for note in operation_notes:
    p=text(note);p.paragraph_format.line_spacing=1.12;p.paragraph_format.space_after=Pt(4)
    for r in p.runs:r.font.size=Pt(10)
md.extend(['## 操作与课前演练','','| 操作 | 按键 |','| --- | --- |'])
md.extend('| '+a+' | '+b+' |' for a,b in operations)
md.extend(['']+operation_notes+[''])
# Greedy cue pagination, conservatively calculated for CJK body text.
def estimate(c):
    h=48
    for s in [c['teacher'] or '（保持安静。）',c['action'],c['hold'],c['reply']]:
        if s:h+=sum(max(1,math.ceil(len(line)/42)) for line in s.split('\n'))*14.2+8
    return h+14
page();used=0;prevsection=None
def silence(c):
    if c['id'].startswith('final-'):return '（保持安静，让画面自行完成。）'
    if c['id'].endswith('-close'):return '（保持安静，等待画面收回并停住。）'
    return '（保持安静，留给学生阅读、讨论或书写。）'
for c in C:
    ht=estimate(c)
    if used and (used+ht>660 or c['section']!=prevsection):page();used=0
    if not used:
        text(c['section']+'  /  教师现场用稿','Caption');used+=28
    title=f"{c['number']:02d}　{c['title']}"
    doc.add_heading(title,2)
    text(f"{c['timing']}　｜　网页定位：{c['id']}",'Caption')
    label('教师说',c['teacher'] or silence(c))
    label('画面',c['action']);label('停留',c['hold']);label('按需承接',c['reply'])
    if c['id']=='final-day':label('推进',c['next'])
    used+=ht;prevsection=c['section']
    md.extend([f'## {title}',f"**{c['timing']} · `{c['id']}`**",'', '**教师说**',c['teacher'] or silence(c),'',f"**画面：** {c['action']}",f"**停留：** {c['hold'] or '完成本节拍后推进。'}"])
    if c['reply']:md.append('**按需承接：** '+c['reply'])
    md.extend(['**推进：** '+c['next'],''])
# Student reading evidence in printable appendix.
page();doc.add_heading('附录一　课堂阅读材料',1)
text('以下编号仅供教师备课定位，不投放到学生主画面。独立摘录之间留空行；出现省略处时，不将前后文字拼成连续全文。标点按课堂显示统一。')
refs=[('hurt','受伤之后','第2章；PDF第34页'),('warmth','外祖母的爱','第1章；PDF第18页'),('window','与表哥一起坐在窗前','第2章；PDF第31页；同伴为米哈伊尔舅舅的儿子萨沙；两块独立摘录'),('grandfather','外祖父的经历与解释','第2章；PDF第36页；各块为独立摘录'),('protect','伸出来的胳膊','第2章；PDF第38页；不可提前揭示后半问答'),('again','保护之后的问答','第2章；PDF第38页'),('anne_context','安妮藏身生活的背景','依据安妮之家资料整理的背景阅读，不是日记原文'),('anne','安妮的安静与愿望','三处独立短引的课堂工作译文，分别来自不同篇章；最后的愿望句为1943年12月24日日记。具体日期与来源见附录二。'),('author','作者的回望','第12章；PDF第243—244页；各块为独立摘录')]
for key,title,ref in refs:
    doc.add_heading(title,2);p=text(ref,'Caption');p.paragraph_format.keep_with_next=True
    for i,q in enumerate(L['passages'][key]):
        p=text(q);p.paragraph_format.keep_with_next=i<len(L['passages'][key])-1
md.extend(['## 附录：阅读证据',''])
for key,title,ref in refs:
    md.extend(['### '+title,ref,'']+L['passages'][key]+[''])
page();doc.add_heading('附录二　材料来源与使用说明',1)
text('本课使用历史作品与历史报道，不把报道中的年龄、环境表述为2026年的即时近况。照片保持原貌，不生成式改画儿童、不添加不存在的环境。出处留在教师资料中，学生投影保持阅读主体清楚。')
source_texts=[
('小说底本','高尔基著，刘辽逸译，《童年》，人民文学出版社。用户提供的281页PDF。这里页码均为文件页序，非纸本页码。'),
('原教案与旧逐字稿','仓库根目录：童年公开课完整教案.docx、teacher_script.docx。既有教案的三层问题、原著摘录及旧稿中的安妮短引与背景仍为本课依据；线条叙事版按57个视觉节拍重排课堂。'),
('列宾原画',f"伊利亚·列宾，《伏尔加河上的纤夫》，1870—1873年，布面油画，俄罗斯博物馆藏，Ж-4056。画作只是劳动处境的视觉参照，不是小说事件的照片。\n{L['sources']['haulers']['source_page']}"),
('安妮的日记与短文',f"安妮之家（Anne Frank House）。三处中文短引沿用旧教师逐字稿中的课堂工作译文，以独立段落呈现，不拼成一篇日记，不冒充已出版中文译本。\n"+'\n'.join(f"第{i+1}处：{e['source_title']}。{e.get('date_note', '日记日期：'+(e.get('diary_entry_date') or '旧稿未标明'))}\n{e['source_url']}" for i,e in enumerate(L['sources']['anne']['excerpts']))+f"\n帮助者背景：{L['sources']['anne']['helpers_source']}"),
('亚门的历史报道',f"Jane Courtney / UNICEF：Hope amid the rubble – Restoring learning in the Gaza Strip，2024年9月10日。图为亚门与父亲；文字改述母亲在家教学、孩子的劳动负担与返校愿望。来源未给年龄，不补写。图片署名：UNICEF 2024。\n{L['sources']['yamen']['source_url']}"),
('瓦尔瓦拉的历史报道',f"UNICEF：As winter bites, needs intensify for families in Ukraine，2022年12月19日。报道时七岁；原图为母女阅读，烛光及停电寒冷的情境来自报道文字，不是给照片增添的事实。图片署名：UNICEF。\n{L['sources']['varvara']['source_url']}"),
('使用与维护','课堂展示不等同于取得素材的任意再分发许可。公开传播、商业使用等须另核对原权利人的条款。本包只保留用户指定的三张既有素材；引用范围和来源不扩大。更改任何节拍后，运行 tools/build_script.py 重新生成Word，并重新检查版式。')]
for title,body in source_texts:
    doc.add_heading(title,3)
    source_lines=body.split('\n')
    for i,line in enumerate(source_lines):
        p=text(line)
        p.paragraph_format.space_after=Pt(0 if i<len(source_lines)-1 else 5)
        # Keep URLs intact across page boundaries and attached to their citation.
        if 'https://' in line:p.paragraph_format.keep_together=True
        if i<len(source_lines)-1 and source_lines[i+1].startswith('https://'):
            p.paragraph_format.keep_with_next=True
    md.extend(['### '+title,body,''])
text('同步标识','Heading 3');text('课堂数据 SHA-256：\n'+SHA,'Caption')
# Narrow package: no embedded fonts, printer settings or thumbnails.
# Put explicit page boundaries on the next real paragraph. An empty break
# paragraph after a full page can otherwise produce an extra blank page in LO.
for p in list(doc.paragraphs):
    if not p.text and p._p.xpath('.//w:br[@w:type="page"]'):
        following=p._p.getnext()
        if following is not None and following.tag==qn('w:p'):
            pr=following.find(qn('w:pPr'))
            if pr is None:pr=OxmlElement('w:pPr');following.insert(0,pr)
            if pr.find(qn('w:pageBreakBefore')) is None:pr.append(OxmlElement('w:pageBreakBefore'))
            p._p.getparent().remove(p._p)
doc.save(DOC)
with zipfile.ZipFile(DOC) as z:
    entries={i.filename:z.read(i.filename) for i in z.infolist()}
# Remove optional stylesWithEffects copy and thumbnail to keep source-control binary compact.
for name in ['word/stylesWithEffects.xml','docProps/thumbnail.jpeg']:
    entries.pop(name,None)
for name in ['word/_rels/document.xml.rels','_rels/.rels','[Content_Types].xml']:
    if name in entries:
        from lxml import etree
        rt=etree.fromstring(entries[name])
        for e in list(rt):
            if any(v in (e.get('Target','')+e.get('PartName','')) for v in ['stylesWithEffects.xml','thumbnail.jpeg']):rt.remove(e)
        entries[name]=etree.tostring(rt,xml_declaration=True,encoding='UTF-8',standalone=True)
# Drop the many unused built-in style definitions, retaining actual used styles + their bases.
from lxml import etree
styles=etree.fromstring(entries['word/styles.xml']);ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
used_styles={'Normal','DefaultParagraphFont','TableNormal','NoList'}
for fn,raw in entries.items():
    if fn.startswith('word/') and fn.endswith('.xml') and fn!='word/styles.xml':
        rt=etree.fromstring(raw)
        for n in rt.xpath('//w:pStyle|//w:rStyle|//w:tblStyle',namespaces=ns):used_styles.add(n.get(qn('w:val')))
byid={x.get(qn('w:styleId')):x for x in styles.findall('w:style',ns)}
for _ in range(5):
    for sid in list(used_styles):
        e=byid.get(sid)
        if e is not None:
            for n in e.xpath('./w:basedOn|./w:next|./w:link',namespaces=ns):used_styles.add(n.get(qn('w:val')))
for sid,e in byid.items():
    if sid not in used_styles:styles.remove(e)
entries['word/styles.xml']=etree.tostring(styles,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(DOC,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
    for name,data in entries.items():z.writestr(name,data)
MD.write_text('\n'.join(md)+'\n',encoding='utf-8')
print(DOC, DOC.stat().st_size, 'bytes;',len(C),'cues;',sum(c['seconds'] for c in C),'seconds')
# Exact text round-trip gate; conditional teacher responses remain separate from spoken script.
actual='\n'.join(p.text for p in Document(DOC).paragraphs)
assert all(c['teacher'] in actual for c in C)
assert all(c['id'] in actual for c in C)
assert not any(name.startswith('word/fonts/') for name in entries)

